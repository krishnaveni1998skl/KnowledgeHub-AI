"""
Document text extraction module using PyMuPDF (fitz) and python-docx.
Rule 4: Detects whether a PDF already contains extractable text before invoking OCR.
Rule 7: Preserves page numbers for source citations.
"""
import io
import re
import logging
from pathlib import Path
from typing import List, Dict, Any
import pymupdf as fitz
from PIL import Image
import docx

from app.services.ocr import is_tesseract_available, ocr_image

logger = logging.getLogger(__name__)

def is_text_sufficient(text: str, min_words: int = 15) -> bool:
    """
    Checks if a page has sufficient extractable text.
    If words < min_words, the page is likely scanned or contains mostly images.
    """
    if not text:
        return False
    words = text.split()
    return len(words) >= min_words

def extract_text_from_pdf(pdf_path: str, min_words: int = 15) -> Dict[str, Any]:
    """
    Extracts text page-by-page from a PDF using PyMuPDF.
    If a page has insufficient text, triggers Tesseract OCR fallback if installed.
    Preserves 1-indexed page numbers for future RAG citations.
    """
    path = Path(pdf_path)
    if not path.is_file():
        raise FileNotFoundError(f"PDF file not found at: {pdf_path}")

    doc = None
    try:
        doc = fitz.open(str(path))
    except Exception as e:
        clean_err = re.sub(r"Failed to open file '.*?' as type", "Invalid or corrupted", str(e))
        clean_err = re.sub(r"[A-Za-z]:\\[^ \t\n\r\f\v'\"]+", "[file]", clean_err)
        raise ValueError(f"Failed to read PDF document (corrupted or invalid format): {clean_err}")

    total_pages = len(doc)
    if total_pages == 0:
        doc.close()
        raise ValueError("PDF document contains 0 pages.")

    pages_result: List[Dict[str, Any]] = []
    used_direct = False
    used_ocr = False

    try:
        for i, page in enumerate(doc):
            page_num = i + 1
            direct_text = page.get_text("text").strip()

            if is_text_sufficient(direct_text, min_words=min_words):
                # Page has digital text layer
                pages_result.append({
                    "page_number": page_num,
                    "text": direct_text,
                    "extraction_type": "direct"
                })
                used_direct = True
            else:
                # Page is image-heavy or scanned; attempt OCR fallback
                if is_tesseract_available():
                    try:
                        logger.info(f"Page {page_num} has insufficient text ({len(direct_text.split())} words). Running Tesseract OCR...")
                        pix = page.get_pixmap(dpi=150)
                        img = Image.open(io.BytesIO(pix.tobytes("png")))
                        ocr_text = ocr_image(img).strip()

                        if ocr_text:
                            pages_result.append({
                                "page_number": page_num,
                                "text": ocr_text,
                                "extraction_type": "ocr"
                            })
                            used_ocr = True
                        else:
                            # OCR returned empty, keep whatever direct text existed
                            pages_result.append({
                                "page_number": page_num,
                                "text": direct_text,
                                "extraction_type": "direct"
                            })
                            if direct_text:
                                used_direct = True
                    except Exception as ocr_err:
                        logger.warning(f"OCR failed on page {page_num}: {ocr_err}. Falling back to direct text.")
                        pages_result.append({
                            "page_number": page_num,
                            "text": direct_text,
                            "extraction_type": "direct"
                        })
                        if direct_text:
                            used_direct = True
                else:
                    # Tesseract not installed, use whatever text PyMuPDF found
                    pages_result.append({
                        "page_number": page_num,
                        "text": direct_text,
                        "extraction_type": "direct"
                    })
                    if direct_text:
                        used_direct = True
    finally:
        doc.close()

    if used_direct and used_ocr:
        overall_method = "hybrid"
    elif used_ocr:
        overall_method = "ocr"
    else:
        overall_method = "direct"

    return {
        "pages": pages_result,
        "page_count": total_pages,
        "extraction_method": overall_method
    }

def extract_text_from_docx(docx_path: str) -> Dict[str, Any]:
    """
    Extracts text and table content from DOCX files using python-docx.
    Gracefully handles formatting and table structures.
    """
    path = Path(docx_path)
    if not path.is_file():
        raise FileNotFoundError(f"DOCX file not found at: {docx_path}")

    try:
        doc = docx.Document(str(path))
    except Exception as e:
        clean_err = re.sub(r"[A-Za-z]:\\[^ \t\n\r\f\v'\"]+", "[file]", str(e))
        raise ValueError(f"Failed to read DOCX file (corrupted or invalid format): {clean_err}")

    content_blocks: List[str] = []

    # 1. Paragraphs
    for para in doc.paragraphs:
        p_text = para.text.strip()
        if p_text:
            content_blocks.append(p_text)

    # 2. Tables
    for table in doc.tables:
        table_rows: List[str] = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            content_blocks.append("\n".join(table_rows))

    combined_text = "\n\n".join(content_blocks).strip()
    if not combined_text:
        raise ValueError("DOCX document contains no readable text.")

    return {
        "pages": [
            {
                "page_number": 1,
                "text": combined_text,
                "extraction_type": "direct"
            }
        ],
        "page_count": 1,
        "extraction_method": "direct"
    }
