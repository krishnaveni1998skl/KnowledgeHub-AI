import io
import os
import tempfile
import uuid
import pytest
import pymupdf as fitz
import docx
from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal
from app.models.schema import Document, DocumentChunk
from app.services.chunker import clean_text, chunk_document, approximate_token_count
from app.services.extractor import extract_text_from_pdf, extract_text_from_docx, is_text_sufficient

def create_sample_pdf(num_pages: int = 2, salt: str = "") -> bytes:
    """Helper to generate a clean, real PDF in memory using PyMuPDF."""
    doc = fitz.open()
    for i in range(num_pages):
        page = doc.new_page()
        text = (
            f"KnowledgeHub AI Section {i + 1} {salt}. "
            f"This is an intelligent corporate documentation system using advanced Retrieval-Augmented Generation. "
            f"All information on page {i + 1} must be accurately cited and indexed for semantic search."
        )
        page.insert_text((50, 72), text, fontsize=12)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes

def create_sample_docx(salt: str = "") -> bytes:
    """Helper to generate a clean, real DOCX document in memory using python-docx."""
    doc = docx.Document()
    doc.add_heading(f"Company Policy & HR Guidelines {salt}", level=1)
    doc.add_paragraph("This document outlines core policies for employees and operations.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Policy Category"
    table.rows[0].cells[1].text = "Requirement Level"
    table.rows[1].cells[0].text = "Information Security"
    table.rows[1].cells[1].text = "Mandatory Compliance"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

# ==========================================
# 1. Chunker & Text Cleaning Tests
# ==========================================

def test_clean_text():
    """Verify clean_text strips control characters, normalizes whitespace, and preserves paragraph breaks."""
    raw = "Hello\x00 World!\xa0\r\n\r\n\r\nThis is a   test.\n\n\n\nNext line."
    cleaned = clean_text(raw)
    assert "\x00" not in cleaned
    assert "\xa0" not in cleaned
    assert "This is a test." in cleaned
    assert "\n\n\n" not in cleaned # Max 1 empty line between paragraphs

def test_chunk_document_preserves_page_numbers():
    """Verify chunk_document retains page citations across multiple pages."""
    pages = [
        {"page_number": 1, "text": "Page one paragraph text about artificial intelligence and retrieval systems."},
        {"page_number": 2, "text": "Page two text with security details and compliance specifications."}
    ]
    chunks = chunk_document(pages, chunk_size=50, chunk_overlap=10)
    assert len(chunks) >= 2

    # Verify page citation associations
    p1_chunks = [c for c in chunks if c["page_number"] == 1]
    p2_chunks = [c for c in chunks if c["page_number"] == 2]
    assert len(p1_chunks) > 0
    assert len(p2_chunks) > 0

    # Verify sequential chunk indexing
    for i, c in enumerate(chunks):
        assert c["chunk_index"] == i
        assert c["token_count"] > 0
        assert len(c["content"]) > 0

# ==========================================
# 2. PDF & DOCX Extractor Tests
# ==========================================

def test_pdf_extraction_real_pdf():
    """Verify PyMuPDF extracts text page-by-page from real PDF and detects digital text."""
    pdf_data = create_sample_pdf(num_pages=2)
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(pdf_data)
        tmp_path = tmp.name

    try:
        result = extract_text_from_pdf(tmp_path)
        assert result["page_count"] == 2
        assert result["extraction_method"] == "direct"
        assert len(result["pages"]) == 2
        assert result["pages"][0]["page_number"] == 1
        assert "KnowledgeHub AI Section 1" in result["pages"][0]["text"]
        assert result["pages"][1]["page_number"] == 2
        assert "KnowledgeHub AI Section 2" in result["pages"][1]["text"]
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)

def test_docx_extraction_real_docx():
    """Verify python-docx extracts paragraphs and tables."""
    docx_data = create_sample_docx()
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(docx_data)
        tmp_path = tmp.name

    try:
        result = extract_text_from_docx(tmp_path)
        assert result["page_count"] == 1
        assert result["extraction_method"] == "direct"
        extracted_text = result["pages"][0]["text"]
        assert "Company Policy & HR Guidelines" in extracted_text
        assert "Information Security" in extracted_text
        assert "Mandatory Compliance" in extracted_text
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)

def test_is_text_sufficient():
    """Verify detection of text-sparse pages to prevent unnecessary OCR."""
    sparse_text = "Img 1"
    rich_text = "This is a full document page with more than fifteen distinct English words describing enterprise knowledge."
    assert is_text_sufficient(sparse_text, min_words=15) is False
    assert is_text_sufficient(rich_text, min_words=15) is True

# ==========================================
# 3. Document API & Upload Validation Tests
# ==========================================

def test_upload_validation_invalid_extension():
    """Reject unsupported file extensions (e.g., .exe, .sh)."""
    with TestClient(app) as client:
        response = client.post(
            "/api/documents/upload",
            files={"file": ("malicious.exe", b"binarycontent", "application/x-msdownload")}
        )
        assert response.status_code == 400
        assert "Unsupported file extension" in response.json()["detail"]

def test_upload_validation_empty_file():
    """Reject 0-byte empty files."""
    with TestClient(app) as client:
        response = client.post(
            "/api/documents/upload",
            files={"file": ("empty.pdf", b"", "application/pdf")}
        )
        assert response.status_code == 400
        assert "empty" in response.json()["detail"].lower()

def test_upload_path_traversal_prevention():
    """Ensure path traversal attempts in filename are sanitized."""
    pdf_bytes = create_sample_pdf(num_pages=1, salt=uuid.uuid4().hex)
    with TestClient(app) as client:
        response = client.post(
            "/api/documents/upload",
            files={"file": ("../../../../evil_traversal.pdf", pdf_bytes, "application/pdf")}
        )
        assert response.status_code == 200
        data = response.json()
        assert ".." not in data["filename"]
        assert data["filename"] == "evil_traversal.pdf"
        # Cleanup uploaded test doc
        client.delete(f"/api/documents/{data['id']}")

def test_upload_pdf_end_to_end_and_sqlite_status():
    """Upload valid PDF, verify SQLite status is 'indexed', and inspect citation chunks."""
    pdf_bytes = create_sample_pdf(num_pages=2, salt=uuid.uuid4().hex)
    with TestClient(app) as client:
        upload_resp = client.post(
            "/api/documents/upload",
            files={"file": ("hr_handbook.pdf", pdf_bytes, "application/pdf")}
        )
        assert upload_resp.status_code == 200
        data = upload_resp.json()
        doc_id = data["id"]
        try:
            assert data["status"] == "indexed"
            assert data["page_count"] == 2
            assert data["chunk_count"] >= 2
            assert data["extraction_method"] == "direct"

            # Check document retrieval
            get_resp = client.get(f"/api/documents/{doc_id}")
            assert get_resp.status_code == 200
            assert get_resp.json()["status"] == "indexed"

            # Check chunk source citations
            chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
            assert chunks_resp.status_code == 200
            chunks_data = chunks_resp.json()
            assert chunks_data["total_chunks"] == data["chunk_count"]
            # Ensure page numbers 1 and 2 exist in chunks
            pages_in_chunks = {c["page_number"] for c in chunks_data["chunks"]}
            assert 1 in pages_in_chunks
            assert 2 in pages_in_chunks
        finally:
            # Cleanup uploaded test doc
            client.delete(f"/api/documents/{doc_id}")

def test_upload_docx_end_to_end():
    """Upload valid DOCX and verify successful ingestion."""
    docx_bytes = create_sample_docx(salt=uuid.uuid4().hex)
    with TestClient(app) as client:
        upload_resp = client.post(
            "/api/documents/upload",
            files={"file": ("security_policy.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert upload_resp.status_code == 200
        data = upload_resp.json()
        try:
            assert data["status"] == "indexed"
            assert data["page_count"] == 1
            assert data["chunk_count"] >= 1
        finally:
            # Cleanup uploaded test doc
            client.delete(f"/api/documents/{data['id']}")


# ==========================================
# 4. Corrupted File & Delete Flow Regression Tests
# ==========================================

def test_corrupted_pdf_upload_error_handling():
    """Verify uploading a corrupted PDF returns 422, marks document as failed in DB, and does not leak server paths."""
    corrupted_bytes = b"GARBAGE_HEADER_NOT_A_VALID_PDF_12345"
    with TestClient(app) as client:
        resp = client.post(
            "/api/documents/upload",
            files={"file": ("corrupted_sample.pdf", corrupted_bytes, "application/pdf")}
        )
        assert resp.status_code == 422
        detail = resp.json()["detail"]
        assert "corrupted" in detail.lower() or "failed to process" in detail.lower()
        # Ensure server local disk paths are not leaked in error response
        assert "C:\\" not in detail
        assert "KNOWLEDGE AI PROJECT" not in detail

        # Verify SQLite record state
        db = SessionLocal()
        try:
            doc = db.query(Document).filter(Document.filename == "corrupted_sample.pdf").order_by(Document.id.desc()).first()
            assert doc is not None
            assert doc.status == "failed"
            assert doc.chunk_count == 0
            assert "C:\\" not in doc.error_message
            doc_id = doc.id
        finally:
            db.close()

        # Clean up
        del_resp = client.delete(f"/api/documents/{doc_id}")
        assert del_resp.status_code == 200


def test_corrupted_docx_upload_error_handling():
    """Verify uploading a corrupted DOCX returns 422 and marks document as failed in DB."""
    corrupted_bytes = b"GARBAGE_DOCX_NOT_A_ZIP_ARCHIVE"
    with TestClient(app) as client:
        resp = client.post(
            "/api/documents/upload",
            files={"file": ("corrupted_sample.docx", corrupted_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert resp.status_code == 422
        detail = resp.json()["detail"]
        assert "failed to process" in detail.lower() or "corrupted" in detail.lower()
        assert "C:\\" not in detail

        db = SessionLocal()
        try:
            doc = db.query(Document).filter(Document.filename == "corrupted_sample.docx").order_by(Document.id.desc()).first()
            assert doc is not None
            assert doc.status == "failed"
            doc_id = doc.id
        finally:
            db.close()

        # Clean up
        del_resp = client.delete(f"/api/documents/{doc_id}")
        assert del_resp.status_code == 200


def test_consistent_delete_flow_removes_sqlite_chunks_and_file():
    """Verify DELETE /api/documents/{id} cascades deletion to chunks and deletes file from disk."""
    pdf_bytes = create_sample_pdf(num_pages=2, salt=uuid.uuid4().hex)
    with TestClient(app) as client:
        upload_resp = client.post(
            "/api/documents/upload",
            files={"file": ("delete_test.pdf", pdf_bytes, "application/pdf")}
        )
        assert upload_resp.status_code == 200
        data = upload_resp.json()
        doc_id = data["id"]

        # Confirm chunks and file exist
        db = SessionLocal()
        try:
            doc = db.query(Document).filter(Document.id == doc_id).first()
            assert doc is not None
            file_path = doc.file_path
            assert os.path.exists(file_path)
            chunk_count = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).count()
            assert chunk_count > 0
        finally:
            db.close()

        # Delete document
        del_resp = client.delete(f"/api/documents/{doc_id}")
        assert del_resp.status_code == 200

        # Verify complete removal
        db = SessionLocal()
        try:
            deleted_doc = db.query(Document).filter(Document.id == doc_id).first()
            assert deleted_doc is None
            remaining_chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).count()
            assert remaining_chunks == 0
        finally:
            db.close()

        assert not os.path.exists(file_path)


def test_duplicate_content_upload_rejected_with_409():
    """Verify uploading identical file content a second time is rejected with HTTP 409 Conflict."""
    pdf_bytes = create_sample_pdf(num_pages=1, salt=uuid.uuid4().hex)
    with TestClient(app) as client:
        resp1 = client.post(
            "/api/documents/upload",
            files={"file": ("unique_doc_test.pdf", pdf_bytes, "application/pdf")}
        )
        assert resp1.status_code == 200
        id1 = resp1.json()["id"]
        hash1 = resp1.json()["file_hash"]
        assert hash1 is not None

        # Attempt to upload identical content with same filename
        resp2 = client.post(
            "/api/documents/upload",
            files={"file": ("unique_doc_test.pdf", pdf_bytes, "application/pdf")}
        )
        assert resp2.status_code == 409
        assert "duplicate document detected" in resp2.json()["detail"].lower()

        # Attempt to upload identical content with different filename
        resp3 = client.post(
            "/api/documents/upload",
            files={"file": ("renamed_copy.pdf", pdf_bytes, "application/pdf")}
        )
        assert resp3.status_code == 409
        assert "duplicate document detected" in resp3.json()["detail"].lower()

        client.delete(f"/api/documents/{id1}")


def test_different_content_same_filename_allowed():
    """Verify uploading different content with the same filename is accepted (versioning/updates)."""
    pdf_bytes_v1 = create_sample_pdf(num_pages=1, salt=uuid.uuid4().hex)
    # Different content produces a different SHA-256 hash
    doc2 = fitz.open()
    p2 = doc2.new_page()
    p2.insert_text((50, 72), f"Version 2 updated content {uuid.uuid4().hex}", fontsize=12)
    pdf_bytes_v2 = doc2.tobytes()
    doc2.close()

    with TestClient(app) as client:
        resp1 = client.post(
            "/api/documents/upload",
            files={"file": ("policy_update.pdf", pdf_bytes_v1, "application/pdf")}
        )
        assert resp1.status_code == 200
        id1 = resp1.json()["id"]

        resp2 = client.post(
            "/api/documents/upload",
            files={"file": ("policy_update.pdf", pdf_bytes_v2, "application/pdf")}
        )
        assert resp2.status_code == 200
        id2 = resp2.json()["id"]

        try:
            assert id1 != id2
            assert resp1.json()["file_hash"] != resp2.json()["file_hash"]
        finally:
            client.delete(f"/api/documents/{id1}")
            client.delete(f"/api/documents/{id2}")


def test_document_file_endpoint_safe_serving():
    """Verify GET /api/documents/{id}/file serves the uploaded file with correct headers."""
    pdf_bytes = create_sample_pdf(num_pages=1, salt=uuid.uuid4().hex)
    with TestClient(app) as client:
        upload_resp = client.post(
            "/api/documents/upload",
            files={"file": ("download_view_test.pdf", pdf_bytes, "application/pdf")}
        )
        assert upload_resp.status_code == 200
        doc_id = upload_resp.json()["id"]

        try:
            file_resp = client.get(f"/api/documents/{doc_id}/file")
            assert file_resp.status_code == 200
            assert "application/pdf" in file_resp.headers["content-type"]
            assert 'inline; filename="download_view_test.pdf"' in file_resp.headers["content-disposition"]
            assert len(file_resp.content) == len(pdf_bytes)
        finally:
            client.delete(f"/api/documents/{doc_id}")


def test_document_file_endpoint_nonexistent_and_security():
    """Verify GET /api/documents/{id}/file returns 404 for invalid ID."""
    with TestClient(app) as client:
        resp = client.get("/api/documents/9999999/file")
        assert resp.status_code == 404


